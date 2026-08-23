from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from datetime import date

OUTPUT = "PhytoScan_Project_Explanation.docx"
CLASSES = [
    "Apple_cedar_apple_rust", "Apple_healthy", "Apple_scab", "Cherry_healthy",
    "Corn_common_rust", "Corn_gray_leaf_spot", "Corn_northern_leaf_blight",
    "Grape_black_rot", "Grape_healthy", "Peach_healthy",
    "Pepper_bell_bacterial_spot", "Pepper_bell_healthy", "Potato_early_blight",
    "Potato_late_blight", "Soybean_healthy", "Tomato_bacterial_spot",
    "Tomato_early_blight", "Tomato_healthy", "Tomato_late_blight",
    "Tomato_leaf_mold", "Tomato_mosaic_virus", "Tomato_septoria_leaf_spot",
    "Tomato_yellow_leaf_curl_virus",
]


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), fill)
    properties.append(element)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, (255, 255, 255))
        shade(table.rows[0].cells[index], '176B5B')
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if row_index % 2 == 1:
                shade(cells[index], 'EAF4F0')
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return table


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    paragraph.add_run(text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style='List Number')
    paragraph.add_run(text)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Heading 1'].font.name = 'Aptos Display'
styles['Heading 1'].font.color.rgb = RGBColor(23, 107, 91)
styles['Heading 2'].font.name = 'Aptos Display'
styles['Heading 2'].font.color.rgb = RGBColor(37, 89, 77)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('PhytoScan')
r.bold = True
r.font.size = Pt(30)
r.font.color.rgb = RGBColor(23, 107, 91)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Plant Disease Classification System')
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = RGBColor(45, 55, 50)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Technical Project Explanation and Demonstration Notes')
r.italic = True
r.font.size = Pt(12)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(f'Prepared on {date.today().isoformat()} | SIH 2026').font.size = Pt(10)
doc.add_paragraph()

p = doc.add_paragraph()
p.add_run('Executive Summary').bold = True
p.runs[0].font.size = Pt(15)
p.runs[0].font.color.rgb = RGBColor(23, 107, 91)
doc.add_paragraph(
    'PhytoScan is a web application that identifies plant diseases from a photograph of a leaf. '
    'A React and Vite frontend accepts an image and sends it to a FastAPI backend. The backend '
    'preprocesses the image and runs a fine-tuned Swin Transformer-S (Swin-S) model in PyTorch. '
    'The service returns the three most likely classes, confidence percentages, plant name, severity, '
    'a short explanation, and recommended remedies.'
)

h = doc.add_heading('1. System Architecture', level=1)
doc.add_paragraph('The project is split into two independently runnable services:')
add_table(doc, ['Layer', 'Technology', 'Responsibility'], [
    ['Frontend', 'React 18 + Vite', 'Image selection, preview, loading state, error handling, and results display.'],
    ['Backend API', 'FastAPI + Uvicorn', 'Validates uploads, exposes health and prediction endpoints, and coordinates inference.'],
    ['Inference engine', 'PyTorch + torchvision', 'Builds Swin-S, loads checkpoint weights, preprocesses images, and calculates probabilities.'],
    ['Metadata layer', 'class_names.py', 'Maintains the model class order and supplies display labels, severity, descriptions, and remedies.'],
], [1.3, 1.5, 4.7])
doc.add_paragraph('Request flow: Browser -> Vite proxy (/api) -> FastAPI -> image preprocessing -> Swin-S -> softmax/top-k -> JSON response -> React result cards.')

h = doc.add_heading('2. Model and Class Configuration', level=1)
doc.add_paragraph(
    'The active checkpoint is swin_model/best_swin_s_vast.pth. It contains a Swin-S state dictionary '
    'with a classifier producing 23 logits. The checkpoint was produced using torch.compile and a training '
    'implementation whose parameter names differ from torchvision. The backend removes the _orig_mod prefix '
    'and maps patch embedding, transformer stages, MLP, downsampling, and classifier keys to the torchvision Swin-S structure.'
)
doc.add_paragraph('The model is configured with:')
add_bullet(doc, 'Architecture: torchvision Swin Transformer-S (Swin-S).')
add_bullet(doc, 'Input tensor shape after preprocessing: 1 x 3 x 224 x 224.')
add_bullet(doc, 'Classifier input features: 768.')
add_bullet(doc, 'Classifier output: 23 logits, one for each class below.')
add_bullet(doc, 'Runtime device: CUDA GPU when available; otherwise CPU. The verified run used CPU.')
add_bullet(doc, 'Model loading: once at backend startup through a singleton, so later requests reuse the loaded model.')

rows = [(str(i).zfill(2), name) for i, name in enumerate(CLASSES)]
doc.add_heading('Active 23-class label order', level=2)
add_table(doc, ['Index', 'Class name used by the API'], rows, [0.8, 5.8])
doc.add_paragraph('The order is important: output index 0 is Apple_cedar_apple_rust, output index 1 is Apple_healthy, and so on. Changing this order would make otherwise correct model predictions display the wrong disease.')

h = doc.add_heading('3. Input and Image Preprocessing', level=1)
doc.add_paragraph('The user supplies one image file through the frontend. The backend accepts:')
add_bullet(doc, 'JPEG/JPG, PNG, or WebP images.')
add_bullet(doc, 'A maximum request file size of 10 MB.')
add_bullet(doc, 'A multipart/form-data field named file.')
add_bullet(doc, 'An image that can be decoded by Pillow; it is converted to RGB.')
doc.add_paragraph('The preprocessing pipeline in backend/model.py is:')
add_number(doc, 'Resize the shorter image dimension to 256 pixels while preserving aspect ratio.')
add_number(doc, 'Center-crop the image to 224 x 224 pixels.')
add_number(doc, 'Convert the image to a PyTorch tensor with channel order RGB.')
add_number(doc, 'Normalize channels with ImageNet mean [0.485, 0.456, 0.406] and standard deviation [0.229, 0.224, 0.225].')
add_number(doc, 'Add a batch dimension, producing a 1 x 3 x 224 x 224 tensor.')
doc.add_paragraph('The application does not store the uploaded image. It reads the bytes for inference and returns the prediction response.')

h = doc.add_heading('4. Inference and Real Timing', level=1)
doc.add_paragraph('During inference, the backend performs the following operations:')
add_number(doc, 'Decode and preprocess the uploaded image.')
add_number(doc, 'Run the tensor through the loaded Swin-S model under torch.no_grad().')
add_number(doc, 'Apply softmax across the 23 logits to obtain class probabilities.')
add_number(doc, 'Select the top three probabilities with torch.topk.')
add_number(doc, 'Attach metadata from class_names.py to each selected class.')
add_table(doc, ['Measurement', 'Observed value', 'Conditions'], [
    ['Real end-to-end prediction request', '2,572 ms (2.57 seconds)', 'Live FastAPI server, CPU device, local PNG sample, multipart HTTP request, model already loaded.'],
    ['Model startup', 'Not included above', 'Checkpoint loading happens during FastAPI startup and is intentionally excluded from per-request timing.'],
    ['Frontend status', 'Successful', 'Vite development server returned HTTP 200 and proxied /api/health successfully.'],
], [2.1, 1.6, 3.0])
doc.add_paragraph(
    'The measured 2.57 seconds is the real CPU result in this development environment and includes HTTP upload, '
    'image decoding, preprocessing, neural-network inference, post-processing, and JSON serialization. The exact time '
    'will vary with CPU/GPU hardware, image size, operating-system load, and whether the model is warm. GPU inference '
    'may be substantially faster. The frontend currently displays “<1s” as a product-facing estimate, but the measured '
    'CPU result should be used for an honest local performance description.'
)

h = doc.add_heading('5. Backend Explanation', level=1)
doc.add_paragraph('backend/main.py defines the HTTP service:')
add_table(doc, ['Endpoint', 'Method', 'Purpose'], [
    ['/', 'GET', 'Confirms that the Plant Disease Classifier API is running.'],
    ['/health', 'GET', 'Reports service status, whether the model is loaded, the absolute checkpoint path, and device.'],
    ['/predict', 'POST', 'Accepts one image file and returns the top-three disease predictions.'],
], [1.2, 0.8, 4.7])
doc.add_paragraph('The /predict endpoint validates the MIME type, rejects empty files, enforces the 10 MB limit, calls predict(), and returns HTTP errors for invalid input, missing models, or inference failures. CORS is enabled for the local React origins.')
doc.add_paragraph('The response is validated with Pydantic models. Each prediction contains:')
add_bullet(doc, 'class_name: machine-readable class identifier in the 23-class order.')
add_bullet(doc, 'display_name: readable disease name.')
add_bullet(doc, 'confidence: rounded probability percentage.')
add_bullet(doc, 'plant: crop associated with the prediction.')
add_bullet(doc, 'severity: None, Moderate, High, Critical, or Unknown depending on metadata.')
add_bullet(doc, 'description: disease explanation when available.')
add_bullet(doc, 'remedies: a list of suggested treatment or management steps.')
doc.add_paragraph('The default backend port is 8001 because port 8000 was occupied by another local service. The backend can also be configured with the MODEL_PATH environment variable.')

h = doc.add_heading('6. Frontend Explanation', level=1)
doc.add_paragraph('frontend/src/App.jsx provides the user workflow:')
add_number(doc, 'Hero section: introduces AI-powered plant diagnostics and links to the scanner.')
add_number(doc, 'Upload zone: supports click-to-browse and drag-and-drop, validates image selection on the client, and shows a local preview.')
add_number(doc, 'Analysis action: sends the selected image as FormData to /api/predict.')
add_number(doc, 'Loading state: disables replacement while inference runs and displays an analysis spinner.')
add_number(doc, 'Results panel: renders the top three predictions returned by the backend.')
add_number(doc, 'Result cards: show the readable name, crop, severity badge, confidence percentage, confidence bar, description, and treatment steps for the top match.')
add_number(doc, 'Reset action: clears the current image, result, and error so another scan can begin.')
add_bullet(doc, 'The Vite development proxy maps /api to http://localhost:8001.')
add_bullet(doc, 'The UI shows supported disease highlights, process steps, feature descriptions, and the 23-class count.')
add_bullet(doc, 'Frontend error handling displays backend validation or connection messages to the user.')
doc.add_paragraph('When a new label does not yet have a detailed entry in DISEASE_INFO, the backend generates readable fallback metadata and a general agricultural-extension-service recommendation. This keeps the response schema complete for all 23 active classes.')

h = doc.add_heading('7. Running the Project', level=1)
doc.add_paragraph('Backend:')
code = doc.add_paragraph()
code.style = 'No Spacing'
code.add_run('cd backend\npython -m uvicorn main:app --host 0.0.0.0 --port 8001').font.name = 'Consolas'
doc.add_paragraph('Frontend, in a second terminal:')
code = doc.add_paragraph()
code.style = 'No Spacing'
code.add_run('cd frontend\nnpm install\nnpm run dev -- --host 0.0.0.0').font.name = 'Consolas'
doc.add_paragraph('Open http://localhost:5173/. The direct API health check is http://localhost:8001/health. The frontend proxy health check is http://localhost:5173/api/health.')

h = doc.add_heading('8. Validation Performed', level=1)
add_bullet(doc, 'The checkpoint classifier head was verified as shape (23, 768) with 23 bias values.')
add_bullet(doc, 'The converted checkpoint loaded successfully into torchvision Swin-S.')
add_bullet(doc, 'Backend Python modules passed compileall syntax validation.')
add_bullet(doc, 'The React frontend passed npm run build.')
add_bullet(doc, 'The live /health endpoint returned model_loaded=true and device=cpu.')
add_bullet(doc, 'A real /predict request returned HTTP 200 with three predictions and complete response fields.')
add_bullet(doc, 'The Vite frontend and /api proxy both returned HTTP 200.')

h = doc.add_heading('9. Limitations and Responsible Use', level=1)
doc.add_paragraph(
    'The classifier provides an image-based prediction, not a guaranteed agricultural diagnosis. Accuracy depends on '
    'image quality, lighting, leaf visibility, camera angle, disease similarity, and whether the image matches the '
    'training distribution. Confidence is the model probability for the selected classes and should not be interpreted '
    'as certainty. Treatment suggestions should be reviewed with a qualified agricultural expert, especially before applying chemicals.'
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('End of project explanation').italic = True
doc.save(OUTPUT)
print(OUTPUT)
